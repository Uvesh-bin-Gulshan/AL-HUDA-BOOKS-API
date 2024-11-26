from import_export.formats.base_formats import Format
from io import BytesIO
from docx import Document

class DOCX(Format):
    def get_title(self):
        return "docx"

    def get_extension(self):
        return "docx"

    def get_content_type(self):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def export_data(self, dataset, **kwargs):
        """Convert dataset to DOCX format."""
        document = Document()
        # Add a title
        document.add_heading("Exported Data", level=1)

        # Add a table with data
        table = document.add_table(rows=1, cols=len(dataset.headers))
        header_cells = table.rows[0].cells
        for i, header in enumerate(dataset.headers):
            header_cells[i].text = header

        for row in dataset.dict:
            row_cells = table.add_row().cells
            for i, key in enumerate(row.keys()):
                row_cells[i].text = str(row[key])

        # Save DOCX to a buffer
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def import_data(self, dataset, file_stream, **kwargs):
        """Read DOCX file and convert it to dataset."""
        document = Document(file_stream)
        table = document.tables[0]  # Assuming the first table has the data
        rows = table.rows

        # Extract headers from the first row
        headers = [cell.text for cell in rows[0].cells]
        dataset.headers = headers

        # Extract data from subsequent rows
        for row in rows[1:]:
            dataset.append([cell.text for cell in row.cells])

        return dataset
